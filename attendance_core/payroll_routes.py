def register_payroll_routes(app, deps):
    BytesIO = deps["BytesIO"]
    Response = deps["Response"]
    build_effective_payroll_adjustments = deps["build_effective_payroll_adjustments"]
    build_employee_payslip_pdf_bytes = deps["build_employee_payslip_pdf_bytes"]
    build_employee_payslip_pdf_filename = deps["build_employee_payslip_pdf_filename"]
    build_payroll_rows = deps["build_payroll_rows"]
    build_payroll_stats = deps["build_payroll_stats"]
    enrich_admin_payroll_run = deps["enrich_admin_payroll_run"]
    execute_db = deps["execute_db"]
    fetchall = deps["fetchall"]
    fetchone = deps["fetchone"]
    flash = deps["flash"]
    format_currency = deps["format_currency"]
    format_datetime_12h = deps["format_datetime_12h"]
    get_db = deps["get_db"]
    get_department_options = deps["get_department_options"]
    get_employee_options = deps["get_employee_options"]
    get_employee_released_payroll_item = deps["get_employee_released_payroll_item"]
    get_employee_released_payroll_runs = deps["get_employee_released_payroll_runs"]
    get_payroll_adjustments = deps["get_payroll_adjustments"]
    get_payroll_employee_filter_label = deps["get_payroll_employee_filter_label"]
    get_payroll_recurring_rule = deps["get_payroll_recurring_rule"]
    get_payroll_recurring_rules = deps["get_payroll_recurring_rules"]
    get_payroll_run = deps["get_payroll_run"]
    get_payroll_pdf_export_folder = deps["get_payroll_pdf_export_folder"]
    get_payslip_download_request_summary = deps["get_payslip_download_request_summary"]
    get_recent_payroll_runs = deps["get_recent_payroll_runs"]
    get_recent_payslip_download_requests = deps["get_recent_payslip_download_requests"]
    get_user_by_id = deps["get_user_by_id"]
    invalidate_reports_cache = deps["invalidate_reports_cache"]
    log_activity = deps["log_activity"]
    login_required = deps["login_required"]
    normalize_payroll_filters = deps["normalize_payroll_filters"]
    now_str = deps["now_str"]
    parse_iso_date = deps["parse_iso_date"]
    parse_money_value = deps["parse_money_value"]
    payroll_filter_redirect_args = deps["payroll_filter_redirect_args"]
    quote = deps["quote"]
    redirect = deps["redirect"]
    render_template = deps["render_template"]
    request = deps["request"]
    review_payslip_download_request = deps["review_payslip_download_request"]
    save_released_payroll_pdf_copy = deps["save_released_payroll_pdf_copy"]
    save_payroll_run_snapshot = deps["save_payroll_run_snapshot"]
    send_from_directory = deps["send_from_directory"]
    session = deps["session"]
    submit_payslip_download_request = deps["submit_payslip_download_request"]
    url_for = deps["url_for"]

    @app.route("/my-payroll")
    @login_required(role="employee")
    def employee_payroll_history():
        user = get_user_by_id(session["user_id"])
        if not user:
            session.clear()
            flash("Your session expired. Please log in again.", "warning")
            return redirect(url_for("login"))

        payroll_runs = get_employee_released_payroll_runs(user["id"], limit=24)
        return render_template(
            "employee_payroll_history.html",
            user=user,
            payroll_runs=payroll_runs
        )


    @app.route("/my-payroll/<int:payroll_run_id>")
    @login_required(role="employee")
    def employee_payslip(payroll_run_id):
        user = get_user_by_id(session["user_id"])
        if not user:
            session.clear()
            flash("Your session expired. Please log in again.", "warning")
            return redirect(url_for("login"))

        payslip = get_employee_released_payroll_item(user["id"], payroll_run_id)
        if not payslip:
            flash("Released payslip not found for your account.", "warning")
            return redirect(url_for("employee_payroll_history"))

        return render_template(
            "employee_payslip.html",
            user=user,
            payslip=payslip,
            download_request=payslip.get("download_request")
        )


    @app.route("/my-payroll/<int:payroll_run_id>/request-download", methods=["POST"])
    @login_required(role="employee")
    def request_employee_payslip_download(payroll_run_id):
        user = get_user_by_id(session["user_id"])
        if not user:
            session.clear()
            flash("Your session expired. Please log in again.", "warning")
            return redirect(url_for("login"))

        payslip = get_employee_released_payroll_item(user["id"], payroll_run_id)
        if not payslip:
            flash("Released payslip not found for your account.", "warning")
            return redirect(url_for("employee_payroll_history"))

        request_row, action = submit_payslip_download_request(user, payslip)
        if action == "approved":
            flash("This payslip PDF is already approved for download.", "success")
        elif action == "pending":
            flash("Your payslip PDF request is already pending payroll admin approval.", "info")
        elif action == "resubmitted":
            flash("Payslip PDF approval was requested again and sent back to payroll admin.", "success")
        else:
            flash("Payslip PDF approval request sent to payroll admin.", "success")

        if request.form.get("from_page") == "history":
            return redirect(url_for("employee_payroll_history"))
        return redirect(url_for("employee_payslip", payroll_run_id=payroll_run_id))


    @app.route("/my-payroll/<int:payroll_run_id>/download.pdf")
    @login_required(role="employee")
    def download_employee_payslip_pdf(payroll_run_id):
        user = get_user_by_id(session["user_id"])
        if not user:
            session.clear()
            flash("Your session expired. Please log in again.", "warning")
            return redirect(url_for("login"))

        payslip = get_employee_released_payroll_item(user["id"], payroll_run_id)
        if not payslip:
            flash("Released payslip not found for your account.", "warning")
            return redirect(url_for("employee_payroll_history"))

        download_request = payslip.get("download_request")
        if not download_request or not download_request.get("is_approved"):
            request_row, action = submit_payslip_download_request(user, payslip)
            if action == "pending":
                flash("Your payslip PDF request is still pending payroll admin approval.", "info")
            elif action == "approved":
                download_request = request_row
            elif action == "resubmitted":
                flash("Your previous request was reopened and sent back to payroll admin for approval.", "warning")
            else:
                flash("You need payroll admin approval before downloading this payslip. A request was sent now.", "warning")
            if not download_request or not download_request.get("is_approved"):
                return redirect(url_for("employee_payslip", payroll_run_id=payroll_run_id))

        pdf_bytes = build_employee_payslip_pdf_bytes(
            payslip,
            printed_at_text=format_datetime_12h(now_str()),
        )
        filename = build_employee_payslip_pdf_filename(payslip)
        return Response(
            pdf_bytes,
            mimetype="application/pdf",
            headers={
                "Content-Disposition": f"attachment; filename=\"{filename}\"; filename*=UTF-8''{quote(filename)}",
                "Cache-Control": "no-store",
            },
        )


    @app.route("/admin/payroll")
    @login_required(role="admin")
    def admin_payroll():
        filters = normalize_payroll_filters(
            request.args.get("period", "this_month"),
            request.args.get("date_from", "").strip(),
            request.args.get("date_to", "").strip(),
            request.args.get("department", "").strip(),
            request.args.get("employee_id", "").strip()
        )
        period = filters["period"]
        department_filter = filters["department_filter"]
        employee_filter = filters["employee_filter"]
        date_from = filters["date_from"]
        date_to = filters["date_to"]
        departments = get_department_options()
        employees = get_employee_options()
        payroll_rows = build_payroll_rows(
            date_from,
            date_to,
            department_filter=department_filter,
            employee_filter=employee_filter
        )
        stats = build_payroll_stats(payroll_rows)
        adjustments = get_payroll_adjustments(date_from, date_to, department_filter=department_filter, employee_filter=employee_filter)
        recurring_rules = get_payroll_recurring_rules(
            department_filter=department_filter,
            employee_filter=employee_filter,
            include_inactive=True
        )
        current_run = get_payroll_run(date_from, date_to, department_filter=department_filter, employee_filter=employee_filter)
        if current_run:
            current_run = enrich_admin_payroll_run(current_run)
        recent_runs = get_recent_payroll_runs()
        editing_recurring_rule = get_payroll_recurring_rule(request.args.get("edit_recurring_rule", ""))

        return render_template(
            "admin_payroll.html",
            payroll_rows=payroll_rows,
            departments=departments,
            employees=employees,
            department_filter=department_filter,
            employee_filter=employee_filter,
            period=period,
            date_from=date_from.strftime("%Y-%m-%d"),
            date_to=date_to.strftime("%Y-%m-%d"),
            stats=stats,
            adjustments=adjustments,
            recurring_rules=recurring_rules,
            current_run=current_run,
            recent_runs=recent_runs,
            employee_filter_label=get_payroll_employee_filter_label(employee_filter),
            editing_recurring_rule=editing_recurring_rule
        )


    @app.route("/admin/payroll/download-requests-panel")
    @login_required(role="admin")
    def admin_payroll_download_requests_panel():
        filters = normalize_payroll_filters(
            request.args.get("period", "this_month"),
            request.args.get("date_from", "").strip(),
            request.args.get("date_to", "").strip(),
            request.args.get("department", "").strip(),
            request.args.get("employee_id", "").strip()
        )
        return render_template(
            "_payroll_download_requests_panel.html",
            payslip_download_requests=get_recent_payslip_download_requests(),
            payslip_request_summary=get_payslip_download_request_summary(),
            period=filters["period"],
            date_from=filters["date_from"].strftime("%Y-%m-%d"),
            date_to=filters["date_to"].strftime("%Y-%m-%d"),
            department_filter=filters["department_filter"],
            employee_filter=filters["employee_filter"],
        )


    @app.route("/admin/payroll/download-requests/<int:request_id>/review", methods=["POST"])
    @login_required(role="admin")
    def review_payslip_download_request_route(request_id):
        redirect_args = payroll_filter_redirect_args(request.form)
        decision = (request.form.get("decision", "") or "").strip().lower()
        admin_note = request.form.get("admin_note", "")

        if decision not in {"approve", "reject"}:
            flash("Choose approve or reject for the payslip PDF request.", "danger")
            return redirect(url_for("admin_payroll", **redirect_args) + "#payroll-download-requests")

        updated_request = review_payslip_download_request(
            request_id,
            session["user_id"],
            decision,
            admin_note=admin_note,
        )
        if not updated_request:
            flash("Payslip download request was not found.", "warning")
            return redirect(url_for("admin_payroll", **redirect_args) + "#payroll-download-requests")

        decision_label = "approved" if decision == "approve" else "rejected"
        flash(
            f"Payslip PDF request for {updated_request.get('employee_name')} was {decision_label}.",
            "success" if decision == "approve" else "info",
        )
        return redirect(url_for("admin_payroll", **redirect_args) + "#payroll-download-requests")


    @app.route("/admin/payroll/adjustments", methods=["POST"])
    @login_required(role="admin")
    def add_payroll_adjustment():
        redirect_args = payroll_filter_redirect_args(request.form)
        filters = normalize_payroll_filters(
            request.form.get("period", "this_month"),
            request.form.get("date_from", ""),
            request.form.get("date_to", ""),
            request.form.get("department", ""),
            request.form.get("employee_id", "")
        )
        user_id_raw = (request.form.get("user_id", "") or "").strip()
        adjustment_type = (request.form.get("adjustment_type", "") or "").strip()
        label = (request.form.get("label", "") or "").strip()
        notes = (request.form.get("notes", "") or "").strip()
        amount = parse_money_value(request.form.get("amount", "0"))

        if not user_id_raw.isdigit():
            flash("Please choose an employee for the payroll adjustment.", "danger")
            return redirect(url_for("admin_payroll", **redirect_args))

        if adjustment_type not in {"Allowance", "Deduction"}:
            flash("Payroll adjustment type must be Allowance or Deduction.", "danger")
            return redirect(url_for("admin_payroll", **redirect_args))

        if not label:
            flash("Payroll adjustment label is required.", "danger")
            return redirect(url_for("admin_payroll", **redirect_args))

        if amount <= 0:
            flash("Payroll adjustment amount must be greater than zero.", "danger")
            return redirect(url_for("admin_payroll", **redirect_args))

        employee = fetchone("""
            SELECT id, full_name, department
            FROM users
            WHERE id = ? AND role = 'employee'
        """, (int(user_id_raw),))
        if not employee:
            flash("Selected employee was not found.", "danger")
            return redirect(url_for("admin_payroll", **redirect_args))

        execute_db("""
            INSERT INTO payroll_adjustments (
                user_id, date_from, date_to, adjustment_type, label,
                amount, notes, created_by, created_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            employee["id"],
            filters["date_from_text"],
            filters["date_to_text"],
            adjustment_type,
            label,
            amount,
            notes or None,
            session["user_id"],
            now_str()
        ), commit=True)

        log_activity(
            session["user_id"],
            "ADD PAYROLL ADJUSTMENT",
            f"{adjustment_type} {format_currency(amount)} for {employee['full_name']} ({label})"
        )
        invalidate_reports_cache()
        flash(
            f"{adjustment_type} added for {employee['full_name']}. Save or re-release this payroll period to update the payslip snapshot.",
            "success"
        )
        return redirect(url_for("admin_payroll", **redirect_args))


    @app.route("/admin/payroll/adjustments/<int:adjustment_id>/delete", methods=["POST"])
    @login_required(role="admin")
    def delete_payroll_adjustment(adjustment_id):
        redirect_args = payroll_filter_redirect_args(request.form)
        adjustment = fetchone("""
            SELECT pa.*, u.full_name AS employee_name
            FROM payroll_adjustments pa
            JOIN users u ON u.id = pa.user_id
            WHERE pa.id = ?
        """, (adjustment_id,))
        if not adjustment:
            flash("Payroll adjustment not found.", "warning")
            return redirect(url_for("admin_payroll", **redirect_args))

        execute_db("DELETE FROM payroll_adjustments WHERE id = ?", (adjustment_id,), commit=True)
        log_activity(
            session["user_id"],
            "DELETE PAYROLL ADJUSTMENT",
            f"Removed {adjustment['adjustment_type']} {format_currency(adjustment['amount'])} for {adjustment['employee_name']} ({adjustment['label']})"
        )
        invalidate_reports_cache()
        flash("Payroll adjustment removed. Save or re-release this payroll period to refresh the payslip snapshot.", "info")
        return redirect(url_for("admin_payroll", **redirect_args))


    @app.route("/admin/payroll/recurring-rules", methods=["POST"])
    @login_required(role="admin")
    def save_payroll_recurring_rule():
        redirect_args = payroll_filter_redirect_args(request.form)
        rule_id_raw = (request.form.get("rule_id", "") or "").strip()
        user_id_raw = (request.form.get("user_id", "") or "").strip()
        adjustment_type = (request.form.get("adjustment_type", "") or "").strip()
        label = (request.form.get("label", "") or "").strip()
        recurrence_type = (request.form.get("recurrence_type", "") or "Every Payroll").strip() or "Every Payroll"
        start_date = (request.form.get("start_date", "") or "").strip()
        end_date = (request.form.get("end_date", "") or "").strip()
        notes = (request.form.get("notes", "") or "").strip()
        amount = parse_money_value(request.form.get("amount", "0"))

        if not user_id_raw.isdigit():
            flash("Please choose an employee for the recurring payroll rule.", "danger")
            return redirect(url_for("admin_payroll", **redirect_args))

        if adjustment_type not in {"Allowance", "Deduction"}:
            flash("Recurring payroll rule type must be Allowance or Deduction.", "danger")
            return redirect(url_for("admin_payroll", **redirect_args))

        if recurrence_type not in {"Every Payroll", "Monthly"}:
            flash("Recurring payroll rule recurrence must be Every Payroll or Monthly.", "danger")
            return redirect(url_for("admin_payroll", **redirect_args))

        if not label:
            flash("Recurring payroll rule label is required.", "danger")
            return redirect(url_for("admin_payroll", **redirect_args))

        if amount <= 0:
            flash("Recurring payroll rule amount must be greater than zero.", "danger")
            return redirect(url_for("admin_payroll", **redirect_args))

        employee = fetchone("""
            SELECT id, full_name
            FROM users
            WHERE id = ? AND role = 'employee'
        """, (int(user_id_raw),))
        if not employee:
            flash("Selected employee was not found.", "danger")
            return redirect(url_for("admin_payroll", **redirect_args))

        start_date_value = parse_iso_date(start_date) if start_date else None
        end_date_value = parse_iso_date(end_date) if end_date else None
        if start_date and not start_date_value:
            flash("Recurring payroll rule start date is invalid.", "danger")
            return redirect(url_for("admin_payroll", **redirect_args))
        if end_date and not end_date_value:
            flash("Recurring payroll rule end date is invalid.", "danger")
            return redirect(url_for("admin_payroll", **redirect_args))
        if start_date_value and end_date_value and start_date_value > end_date_value:
            flash("Recurring payroll rule end date must not be earlier than the start date.", "danger")
            return redirect(url_for("admin_payroll", **redirect_args))
        if recurrence_type == "Monthly" and not start_date_value:
            flash("Monthly recurring rules need a start date so the monthly anchor day is clear.", "danger")
            return redirect(url_for("admin_payroll", **redirect_args))

        timestamp = now_str()
        if rule_id_raw.isdigit():
            existing = fetchone("""
                SELECT *
                FROM payroll_recurring_rules
                WHERE id = ?
            """, (int(rule_id_raw),))
            if not existing:
                flash("Recurring payroll rule not found.", "warning")
                return redirect(url_for("admin_payroll", **redirect_args))

            execute_db("""
                UPDATE payroll_recurring_rules
                SET user_id = ?, adjustment_type = ?, label = ?, amount = ?,
                    recurrence_type = ?, start_date = ?, end_date = ?, notes = ?, updated_at = ?
                WHERE id = ?
            """, (
                employee["id"],
                adjustment_type,
                label,
                amount,
                recurrence_type,
                start_date or None,
                end_date or None,
                notes or None,
                timestamp,
                existing["id"]
            ), commit=True)
            log_activity(
                session["user_id"],
                "UPDATE PAYROLL RECURRING RULE",
                f"Updated {adjustment_type.lower()} rule {label} for {employee['full_name']}."
            )
            invalidate_reports_cache()
            flash("Recurring payroll rule updated.", "success")
        else:
            execute_db("""
                INSERT INTO payroll_recurring_rules (
                    user_id, adjustment_type, label, amount, recurrence_type,
                    start_date, end_date, notes, is_active, created_by, created_at, updated_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                employee["id"],
                adjustment_type,
                label,
                amount,
                recurrence_type,
                start_date or None,
                end_date or None,
                notes or None,
                1,
                session["user_id"],
                timestamp,
                timestamp
            ), commit=True)
            log_activity(
                session["user_id"],
                "ADD PAYROLL RECURRING RULE",
                f"Added {adjustment_type.lower()} recurring rule {label} for {employee['full_name']}."
            )
            invalidate_reports_cache()
            flash("Recurring payroll rule added. It will be included automatically when the rule matches the payroll period.", "success")

        return redirect(url_for("admin_payroll", **redirect_args))


    @app.route("/admin/payroll/recurring-rules/<int:rule_id>/toggle", methods=["POST"])
    @login_required(role="admin")
    def toggle_payroll_recurring_rule(rule_id):
        redirect_args = payroll_filter_redirect_args(request.form)
        rule = fetchone("""
            SELECT prr.*, u.full_name AS employee_name
            FROM payroll_recurring_rules prr
            JOIN users u ON u.id = prr.user_id
            WHERE prr.id = ?
        """, (rule_id,))
        if not rule:
            flash("Recurring payroll rule not found.", "warning")
            return redirect(url_for("admin_payroll", **redirect_args))

        new_status = 0 if int(rule.get("is_active") or 0) == 1 else 1
        execute_db("""
            UPDATE payroll_recurring_rules
            SET is_active = ?, updated_at = ?
            WHERE id = ?
        """, (new_status, now_str(), rule_id), commit=True)
        log_activity(
            session["user_id"],
            "TOGGLE PAYROLL RECURRING RULE",
            f"{'Activated' if new_status == 1 else 'Paused'} recurring rule {rule['label']} for {rule['employee_name']}."
        )
        invalidate_reports_cache()
        flash(
            f"Recurring payroll rule {'activated' if new_status == 1 else 'paused'}.",
            "success" if new_status == 1 else "info"
        )
        return redirect(url_for("admin_payroll", **redirect_args))


    @app.route("/admin/payroll/recurring-rules/<int:rule_id>/delete", methods=["POST"])
    @login_required(role="admin")
    def delete_payroll_recurring_rule(rule_id):
        redirect_args = payroll_filter_redirect_args(request.form)
        rule = fetchone("""
            SELECT prr.*, u.full_name AS employee_name
            FROM payroll_recurring_rules prr
            JOIN users u ON u.id = prr.user_id
            WHERE prr.id = ?
        """, (rule_id,))
        if not rule:
            flash("Recurring payroll rule not found.", "warning")
            return redirect(url_for("admin_payroll", **redirect_args))

        execute_db("DELETE FROM payroll_recurring_rules WHERE id = ?", (rule_id,), commit=True)
        log_activity(
            session["user_id"],
            "DELETE PAYROLL RECURRING RULE",
            f"Deleted recurring rule {rule['label']} for {rule['employee_name']}."
        )
        invalidate_reports_cache()
        flash("Recurring payroll rule deleted.", "info")
        return redirect(url_for("admin_payroll", **redirect_args))


    @app.route("/admin/payroll/run", methods=["POST"])
    @login_required(role="admin")
    def save_payroll_run():
        redirect_args = payroll_filter_redirect_args(request.form)
        filters = normalize_payroll_filters(
            request.form.get("period", "this_month"),
            request.form.get("date_from", ""),
            request.form.get("date_to", ""),
            request.form.get("department", ""),
            request.form.get("employee_id", "")
        )
        action = (request.form.get("run_action", "draft") or "draft").strip().lower()
        status = "Released" if action == "release" else "Draft"
        notes = (request.form.get("notes", "") or "").strip()
        payroll_rows = build_payroll_rows(
            filters["date_from"],
            filters["date_to"],
            department_filter=filters["department_filter"],
            employee_filter=filters["employee_filter"]
        )

        if not payroll_rows:
            flash("There are no payroll rows in this view to save.", "warning")
            return redirect(url_for("admin_payroll", **redirect_args))

        if status == "Released" and any(row["has_rate"] == 0 for row in payroll_rows):
            flash("Release is blocked until every employee in view has an hourly rate.", "danger")
            return redirect(url_for("admin_payroll", **redirect_args))

        try:
            payroll_run, _ = save_payroll_run_snapshot(
                filters["date_from"],
                filters["date_to"],
                department_filter=filters["department_filter"],
                employee_filter=filters["employee_filter"],
                status=status,
                notes=notes,
                actor_id=session["user_id"]
            )
        except ValueError as exc:
            flash(str(exc), "warning")
            return redirect(url_for("admin_payroll", **redirect_args))
        log_activity(
            session["user_id"],
            "RELEASE PAYROLL" if status == "Released" else "SAVE PAYROLL DRAFT",
            f"{status} payroll for {filters['date_from_text']} to {filters['date_to_text']} ({len(payroll_rows)} employee row(s))"
        )
        saved_pdf = None
        if status == "Released":
            saved_pdf = save_released_payroll_pdf_copy(
                payroll_run["id"],
                printed_at_text=format_datetime_12h(now_str()),
            )
        invalidate_reports_cache()
        flash(
            f"Payroll {status.lower()} saved for {filters['date_from_text']} to {filters['date_to_text']}.",
            "success"
        )
        if saved_pdf:
            flash(f"Admin PDF copy saved as {saved_pdf['payroll_pdf_filename']}.", "success")
        return redirect(url_for("admin_payroll", **redirect_args))


    @app.route("/admin/payroll/runs/<int:payroll_run_id>/download.pdf")
    @login_required(role="admin")
    def download_admin_payroll_pdf(payroll_run_id):
        payroll_run = fetchone("SELECT * FROM payroll_runs WHERE id = ?", (payroll_run_id,))
        if not payroll_run or payroll_run["status"] != "Released":
            flash("Released payroll PDF was not found.", "warning")
            return redirect(url_for("admin_payroll", **payroll_filter_redirect_args(request.args)))

        payroll_run = save_released_payroll_pdf_copy(
            payroll_run_id,
            printed_at_text=format_datetime_12h(now_str()),
        )
        filename = payroll_run.get("payroll_pdf_filename") if payroll_run else ""

        if not filename:
            flash("Unable to prepare the payroll PDF copy.", "danger")
            return redirect(url_for("admin_payroll", **payroll_filter_redirect_args(request.args)))

        return send_from_directory(
            get_payroll_pdf_export_folder(),
            filename,
            mimetype="application/pdf",
            as_attachment=True,
            download_name=filename,
        )


    @app.route("/admin/payroll/runs/<int:payroll_run_id>/delete", methods=["POST"])
    @login_required(role="admin")
    def delete_payroll_run(payroll_run_id):
        redirect_args = payroll_filter_redirect_args(request.form)
        payroll_run = fetchone("""
            SELECT pr.*, creator.full_name AS created_by_name
            FROM payroll_runs pr
            LEFT JOIN users creator ON creator.id = pr.created_by
            WHERE pr.id = ?
        """, (payroll_run_id,))

        if not payroll_run:
            flash("Payroll run not found.", "warning")
            return redirect(url_for("admin_payroll", **redirect_args))

        payroll_run = enrich_admin_payroll_run(payroll_run)
        if payroll_run["status"] != "Draft":
            flash("Only draft payroll snapshots can be deleted.", "danger")
            return redirect(url_for("admin_payroll", **redirect_args))

        db = get_db()
        try:
            execute_db("DELETE FROM payroll_run_item_adjustments WHERE payroll_run_id = ?", (payroll_run_id,))
            execute_db("DELETE FROM payroll_run_items WHERE payroll_run_id = ?", (payroll_run_id,))
            execute_db("DELETE FROM payroll_runs WHERE id = ?", (payroll_run_id,))
            db.commit()
        except Exception:
            db.rollback()
            raise

        log_activity(
            session["user_id"],
            "DELETE PAYROLL DRAFT",
            f"Deleted draft payroll snapshot {payroll_run['period_label']} ({payroll_run['item_count']} employee row(s))"
        )
        invalidate_reports_cache()
        flash(f"Deleted draft payroll snapshot for {payroll_run['period_label']}.", "info")
        return redirect(url_for("admin_payroll", **redirect_args))


    @app.route("/admin/payroll/bulk-release", methods=["POST"])
    @login_required(role="admin")
    def bulk_release_payroll_runs():
        redirect_args = payroll_filter_redirect_args(request.form)
        raw_ids = request.form.getlist("payroll_run_ids")
        payroll_run_ids = []
        seen_ids = set()
        for raw_id in raw_ids:
            raw_text = str(raw_id or "").strip()
            if not raw_text.isdigit():
                continue
            run_id = int(raw_text)
            if run_id in seen_ids:
                continue
            seen_ids.add(run_id)
            payroll_run_ids.append(run_id)

        if not payroll_run_ids:
            flash("Select at least one draft payroll run to bulk release.", "warning")
            return redirect(url_for("admin_payroll", **redirect_args))

        placeholders = ", ".join(["?"] * len(payroll_run_ids))
        runs = [
            enrich_admin_payroll_run(row)
            for row in fetchall(f"""
                SELECT pr.*, creator.full_name AS created_by_name
                FROM payroll_runs pr
                LEFT JOIN users creator ON creator.id = pr.created_by
                WHERE pr.id IN ({placeholders})
                ORDER BY pr.id DESC
            """, tuple(payroll_run_ids))
        ]
        run_map = {int(run["id"]): run for run in runs}

        releasable_runs = []
        skipped_messages = []
        for requested_run_id in payroll_run_ids:
            run = run_map.get(requested_run_id)
            if not run:
                skipped_messages.append(f"Run #{requested_run_id} was not found.")
                continue
            if run["status"] == "Released":
                skipped_messages.append(f"{run['period_label']} was already released.")
                continue
            if not run["can_release"]:
                skipped_messages.append(f"{run['period_label']} was skipped: {run['release_block_reason']}")
                continue
            releasable_runs.append(run)

        if not releasable_runs:
            flash("None of the selected payroll runs were eligible for release.", "warning")
            for message in skipped_messages[:4]:
                flash(message, "info")
            return redirect(url_for("admin_payroll", **redirect_args))

        timestamp = now_str()
        db = get_db()
        try:
            for run in releasable_runs:
                execute_db("""
                    UPDATE payroll_runs
                    SET status = 'Released',
                        updated_at = ?,
                        released_at = ?
                    WHERE id = ?
                """, (timestamp, timestamp, run["id"]))
            db.commit()
        except Exception:
            db.rollback()
            raise

        saved_pdf_count = 0
        for run in releasable_runs:
            if save_released_payroll_pdf_copy(
                run["id"],
                printed_at_text=format_datetime_12h(timestamp),
            ):
                saved_pdf_count += 1

        released_labels = ", ".join(run["period_label"] for run in releasable_runs[:3])
        if len(releasable_runs) > 3:
            released_labels += f", and {len(releasable_runs) - 3} more"
        log_activity(
            session["user_id"],
            "BULK RELEASE PAYROLL",
            f"Released {len(releasable_runs)} payroll run(s): {released_labels}"
        )
        invalidate_reports_cache()
        flash(f"Released {len(releasable_runs)} payroll run(s).", "success")
        if saved_pdf_count:
            flash(f"Saved {saved_pdf_count} admin payroll PDF copy/copies.", "success")
        for message in skipped_messages[:4]:
            flash(message, "info")
        return redirect(url_for("admin_payroll", **redirect_args))


    @app.route("/admin/payroll/export.xlsx")
    @login_required(role="admin")
    def export_admin_payroll_excel():
        filters = normalize_payroll_filters(
            request.args.get("period", "this_month"),
            request.args.get("date_from", ""),
            request.args.get("date_to", ""),
            request.args.get("department", ""),
            request.args.get("employee_id", "")
        )
        payroll_rows = build_payroll_rows(
            filters["date_from"],
            filters["date_to"],
            department_filter=filters["department_filter"],
            employee_filter=filters["employee_filter"]
        )
        adjustments = build_effective_payroll_adjustments(
            filters["date_from"],
            filters["date_to"],
            department_filter=filters["department_filter"],
            employee_filter=filters["employee_filter"]
        )
        current_run = get_payroll_run(
            filters["date_from"],
            filters["date_to"],
            department_filter=filters["department_filter"],
            employee_filter=filters["employee_filter"]
        )
        stats = build_payroll_stats(payroll_rows)

        try:
            from openpyxl import Workbook
        except Exception:
            flash("Excel export requires openpyxl. Install dependencies and try again.", "danger")
            return redirect(url_for("admin_payroll", **payroll_filter_redirect_args(request.args)))

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Payroll Summary"
        sheet.append(["Payroll Period", f"{filters['date_from_text']} to {filters['date_to_text']}"])
        sheet.append(["Department Filter", filters["department_filter"] or "All Departments"])
        sheet.append(["Employee Filter", get_payroll_employee_filter_label(filters["employee_filter"])])
        sheet.append(["Run Status", current_run["status"] if current_run else "Not Saved"])
        sheet.append(["Gross Payroll", stats["total_gross"]])
        sheet.append(["Overtime Pay", stats["total_overtime_pay"]])
        sheet.append(["Allowances", stats["total_allowances"]])
        sheet.append(["Deductions", stats["total_deductions"]])
        sheet.append(["Final Payroll", stats["total_final_pay"]])
        sheet.append([])
        sheet.append([
            "Employee", "Username", "Department", "Position", "Hourly Rate",
            "Days Worked", "Total Hours", "Overtime Hours", "Late Minutes",
            "Break Minutes", "Suspension Days", "Lost Pay Estimate",
            "Gross Pay", "Overtime Pay", "Allowances", "Deductions", "Final Pay", "Status"
        ])
        for row in payroll_rows:
            sheet.append([
                row["full_name"],
                row["username"],
                row["department"],
                row["position"],
                row["hourly_rate"],
                row["days_worked"],
                row["total_hours"],
                row["overtime_hours"],
                row["late_minutes"],
                row["break_minutes"],
                row["suspension_days"],
                row["suspension_pay"],
                row["gross_pay"],
                row["overtime_pay"],
                row["allowances"],
                row["deductions"],
                row["final_pay"],
                row["status_label"],
            ])

        adjustment_sheet = workbook.create_sheet(title="Adjustments")
        adjustment_sheet.append(["Employee", "Source", "Recurrence", "Type", "Label", "Amount", "Notes", "Created By", "Created At"])
        for adjustment in adjustments:
            adjustment_sheet.append([
                adjustment["employee_name"],
                adjustment.get("source_kind") or "Manual",
                adjustment.get("recurrence_type") or "",
                adjustment["adjustment_type"],
                adjustment["label"],
                adjustment["amount"],
                adjustment["notes"] or "",
                adjustment["created_by_name"] or "",
                adjustment["created_at"],
            ])

        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return Response(
            output.getvalue(),
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f'attachment; filename="payroll-summary-{filters["date_from_text"]}-to-{filters["date_to_text"]}.xlsx"'
            }
        )


    @app.route("/admin/payroll/print")
    @login_required(role="admin")
    def print_admin_payroll():
        filters = normalize_payroll_filters(
            request.args.get("period", "this_month"),
            request.args.get("date_from", ""),
            request.args.get("date_to", ""),
            request.args.get("department", ""),
            request.args.get("employee_id", "")
        )
        payroll_rows = build_payroll_rows(
            filters["date_from"],
            filters["date_to"],
            department_filter=filters["department_filter"],
            employee_filter=filters["employee_filter"]
        )
        stats = build_payroll_stats(payroll_rows)
        adjustments = build_effective_payroll_adjustments(
            filters["date_from"],
            filters["date_to"],
            department_filter=filters["department_filter"],
            employee_filter=filters["employee_filter"]
        )
        current_run = get_payroll_run(
            filters["date_from"],
            filters["date_to"],
            department_filter=filters["department_filter"],
            employee_filter=filters["employee_filter"]
        )

        return render_template(
            "admin_payroll_print.html",
            payroll_rows=payroll_rows,
            stats=stats,
            adjustments=adjustments,
            current_run=current_run,
            department_filter=filters["department_filter"],
            employee_filter_label=get_payroll_employee_filter_label(filters["employee_filter"]),
            date_from=filters["date_from_text"],
            date_to=filters["date_to_text"],
        )
